#!/usr/bin/env python3
"""Convert a static HTML deck to PPTX.

E2: rasterize the non-text visual layer, then recreate rendered text fragments as editable
PowerPoint text boxes using WeasyPrint's layout tree.
E3: rasterize the full slide for maximum visual fidelity.

This intentionally does not claim E1/native chart parity. Use a native slide tool for specific
slides/elements when E1 is required.
"""
from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt
from weasyprint import HTML

from render_html_deck import print_css, pdf_to_pngs


def color_to_rgb(color) -> RGBColor:
    try:
        coords = color.coordinates
        if getattr(color, "space", "srgb") != "srgb":
            color = color.to("srgb")
            coords = color.coordinates
        vals = [max(0, min(255, round(float(v) * 255))) for v in coords[:3]]
        return RGBColor(*vals)
    except Exception:
        return RGBColor(17, 17, 17)


def first_font(style) -> str:
    try:
        fam = style["font_family"]
        if fam:
            return str(fam[0])
    except Exception:
        pass
    return "Aptos"


def rendered_document(source: Path, width: int, height: int):
    return HTML(filename=str(source), base_url=str(source.parent)).render(
        stylesheets=[print_css(width, height, hide_text=False)]
    )


def _element_text(el, tag: str) -> str:
    semantic = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "th", "blockquote", "figcaption", "button", "label"}
    if tag in semantic:
        text = " ".join("".join(el.itertext()).split())
        if tag == "li" and text:
            text = "• " + text
        return text
    direct = (el.text or "").strip()
    return " ".join(direct.split())


def extract_text_blocks(doc) -> list[list[dict]]:
    pages: list[list[dict]] = []
    preferred_types = {"BlockBox": 3, "InlineBlockBox": 2, "InlineBox": 1}

    def walk_boxes(box):
        actual = getattr(box, "_box", box)
        yield actual
        try:
            children = actual.all_children()
        except Exception:
            children = []
        for child in children:
            yield from walk_boxes(child)

    for page in doc.pages:
        chosen = {}
        for box in walk_boxes(page._page_box):
            el = getattr(box, "element", None)
            if el is None:
                continue
            tag = str(getattr(el, "tag", "")).lower()
            if tag in {"html", "body", "main", "section", "article", "script", "style", "svg", "canvas"}:
                continue
            text = _element_text(el, tag)
            if not text:
                continue
            tname = type(box).__name__
            rank = preferred_types.get(tname, 0)
            if not rank:
                continue
            key = id(el)
            prev = chosen.get(key)
            if prev is not None and prev[0] >= rank:
                continue
            chosen[key] = (rank, box, text, tag)

        items = []
        for _rank, box, text, tag in chosen.values():
            style = box.style
            align = str(style.get("text_align_all", "start"))
            if align in {"start", "left"}:
                align = "left"
            elif align in {"end", "right"}:
                align = "right"
            elif align == "center":
                align = "center"
            elif align == "justify":
                align = "justify"
            else:
                align = "left"
            items.append({
                "text": text,
                "tag": tag,
                "x": float(box.position_x),
                "y": float(box.position_y),
                "w": float(box.width),
                "h": float(box.height),
                "font_size": float(style["font_size"]),
                "font_family": first_font(style),
                "font_weight": int(style["font_weight"]),
                "font_style": str(style["font_style"]),
                "color": color_to_rgb(style["color"]),
                "align": align,
            })
        items.sort(key=lambda x: (x["y"], x["x"]))
        pages.append(items)
    return pages


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", type=Path)
    ap.add_argument("output", type=Path)
    ap.add_argument("--mode", choices=["E2", "E3", "e2", "e3"], default="E2")
    ap.add_argument("--width", type=int, default=1280)
    ap.add_argument("--height", type=int, default=720)
    ns = ap.parse_args()

    source = ns.input.resolve()
    output = ns.output.resolve()
    mode = ns.mode.upper()
    if not source.exists():
        raise SystemExit(f"Input not found: {source}")

    width_in = 13.333333
    height_in = width_in * ns.height / ns.width
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    blank = prs.slide_layouts[6]
    while len(prs.slides):
        slide_id = prs.slides._sldIdLst[-1]
        prs.part.drop_rel(slide_id.rId)
        del prs.slides._sldIdLst[-1]

    doc = rendered_document(source, ns.width, ns.height)
    text_pages = extract_text_blocks(doc) if mode == "E2" else [[] for _ in doc.pages]

    with tempfile.TemporaryDirectory(prefix="ppt-god-export-") as td:
        td_path = Path(td)
        pdf = td_path / "background.pdf"
        HTML(filename=str(source), base_url=str(source.parent)).write_pdf(
            str(pdf), stylesheets=[print_css(ns.width, ns.height, hide_text=(mode == "E2"))]
        )
        img_dir = td_path / "png"
        images = pdf_to_pngs(pdf, img_dir, dpi=96)
        if len(images) != len(doc.pages):
            raise SystemExit(f"Page count mismatch: layout={len(doc.pages)} raster={len(images)}")

        sx = width_in / ns.width
        sy = height_in / ns.height
        for page_index, image in enumerate(images):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
            if mode == "E2":
                for item in text_pages[page_index]:
                    x = max(0.0, item["x"] * sx)
                    y = max(0.0, item["y"] * sy)
                    w = max(0.03, item["w"] * sx * 1.035)
                    h = max(0.03, item["h"] * sy * 1.03)
                    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
                    tf = shape.text_frame
                    tf.clear()
                    tf.margin_left = 0
                    tf.margin_right = 0
                    tf.margin_top = 0
                    tf.margin_bottom = 0
                    tf.word_wrap = False
                    tf.vertical_anchor = MSO_ANCHOR.TOP
                    p = tf.paragraphs[0]
                    from pptx.enum.text import PP_ALIGN
                    p.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER, "justify": PP_ALIGN.JUSTIFY}.get(item.get("align"), PP_ALIGN.LEFT)
                    run = p.add_run()
                    run.text = item["text"]
                    run.font.name = item["font_family"]
                    run.font.size = Pt(max(5.0, min(96.0, item["font_size"] * 0.75)))
                    run.font.bold = item["font_weight"] >= 600
                    run.font.italic = item["font_style"].lower() == "italic"
                    run.font.color.rgb = item["color"]

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    print(f"Wrote {output} ({len(prs.slides)} slides, {mode})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Extract a reusable Brand Pack scaffold from a corporate PPTX.

This script intentionally performs *forensics*, not full template preservation. It extracts
fonts/colors/assets, estimates a safe content area, detects representative slides, and creates
HTML/CSS shells for human/agent refinement.
"""
from __future__ import annotations

import argparse
import collections
import hashlib
import json
import math
import re
import shutil
import subprocess
from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def norm_text(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip())


def color_hex(color_format) -> str | None:
    try:
        rgb = color_format.rgb
        if rgb is not None:
            return str(rgb).upper()
    except Exception:
        pass
    return None


def shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        return norm_text(shape.text)
    return ""


def shape_signature(shape, sw: int, sh: int) -> str:
    def q(v: int, total: int) -> int:
        return round((v / total) * 200) if total else 0
    typ = int(shape.shape_type)
    parts = [typ, q(shape.left, sw), q(shape.top, sh), q(shape.width, sw), q(shape.height, sh)]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            parts.append(hashlib.sha1(shape.image.blob).hexdigest()[:12])
        except Exception:
            pass
    else:
        t = shape_text(shape)
        if t and len(t) <= 60:
            parts.append(t.lower())
    return "|".join(map(str, parts))


def iter_shapes(prs: Presentation) -> Iterable[Any]:
    seen = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            yield ("slide", shape)
        layout = slide.slide_layout
        if id(layout) not in seen:
            seen.add(id(layout))
            for shape in layout.shapes:
                yield ("layout", shape)
        master = layout.slide_master
        if id(master) not in seen:
            seen.add(id(master))
            for shape in master.shapes:
                yield ("master", shape)


def collect_fonts_and_colors(prs: Presentation):
    fonts = collections.Counter()
    colors = collections.Counter()
    for _scope, shape in iter_shapes(prs):
        try:
            fill = shape.fill
            if fill and getattr(fill, "fore_color", None):
                c = color_hex(fill.fore_color)
                if c:
                    colors[c] += 1
        except Exception:
            pass
        try:
            line = shape.line
            if line and getattr(line, "color", None):
                c = color_hex(line.color)
                if c:
                    colors[c] += 1
        except Exception:
            pass
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    c = color_hex(run.font.color)
                    if c:
                        colors[c] += 1
    return fonts, colors


def slide_all_text(slide) -> str:
    chunks = []
    for shape in slide.shapes:
        t = shape_text(shape)
        if t:
            chunks.append(t)
    return "\n".join(chunks)


def detect_representative_slides(prs: Presentation):
    n = len(prs.slides)
    cover = 1 if n else None
    ending = n if n else None
    toc = None
    patterns = ["目录", "contents", "agenda", "table of contents", "议程"]
    for i in range(1, min(n, 6) + 1):
        slide = prs.slides[i - 1]
        text = slide_all_text(slide).lower()
        if any(p in text for p in patterns):
            toc = i
            break

    excluded = {x for x in [cover, ending, toc] if x}
    candidates = [i for i in range(1, n + 1) if i not in excluded]
    if candidates:
        scored = []
        for i in candidates:
            slide = prs.slides[i - 1]
            text_len = len(slide_all_text(slide))
            score = len(slide.shapes) * 4 + min(text_len, 600) / 40
            scored.append((score, i))
        scored.sort(reverse=True)
        body = scored[0][1]
    else:
        body = cover
    section = None
    for i in candidates:
        slide = prs.slides[i - 1]
        txt = slide_all_text(slide)
        if 0 < len(txt) < 80 and len(slide.shapes) <= 8:
            section = i
            break
    return {"cover": cover, "contents": toc, "section": section, "body": body, "ending": ending}


def body_slide_indexes(n: int, reps: dict[str, int | None]) -> list[int]:
    excluded = {v for k, v in reps.items() if v and k in {"cover", "contents", "section", "ending"}}
    body = [i for i in range(1, n + 1) if i not in excluded]
    return body or ([reps["body"]] if reps.get("body") else [])


def common_edge_shapes(prs: Presentation, body_indexes: list[int]):
    sw, sh = prs.slide_width, prs.slide_height
    counts = collections.Counter()
    by_slide = {}
    for idx in body_indexes:
        sigs = []
        for shape in prs.slides[idx - 1].shapes:
            sig = shape_signature(shape, sw, sh)
            sigs.append((sig, shape))
            counts[sig] += 1
        by_slide[idx] = sigs
    threshold = max(2, math.ceil(len(body_indexes) * 0.5)) if len(body_indexes) > 1 else 1
    rep_idx = body_indexes[len(body_indexes) // 2] if body_indexes else 1
    result = []
    for sig, shape in by_slide.get(rep_idx, []):
        left = shape.left / sw
        top = shape.top / sh
        right = (shape.left + shape.width) / sw
        bottom = (shape.top + shape.height) / sh
        full = shape.width >= sw * 0.92 and shape.height >= sh * 0.92
        edge = top < 0.18 or bottom > 0.84 or left < 0.04 or right > 0.96
        if counts[sig] >= threshold and (edge or full):
            result.append({
                "signature": sig,
                "name": shape.name,
                "type": int(shape.shape_type),
                "left": left,
                "top": top,
                "right": right,
                "bottom": bottom,
                "text": shape_text(shape)[:120],
                "repeated_on": counts[sig],
            })
    return result


def estimate_safe_area(common: list[dict[str, Any]]):
    x0, y0, x1, y1 = 0.055, 0.12, 0.945, 0.88
    for s in common:
        if (s["right"] - s["left"]) > 0.9 and (s["bottom"] - s["top"]) > 0.9:
            continue
        if s["top"] < 0.25 and s["bottom"] < 0.38:
            y0 = max(y0, min(0.28, s["bottom"] + 0.018))
        if s["bottom"] > 0.75 and s["top"] > 0.60:
            y1 = min(y1, max(0.72, s["top"] - 0.018))
        if s["left"] < 0.12 and s["right"] < 0.28:
            x0 = max(x0, min(0.18, s["right"] + 0.012))
        if s["right"] > 0.88 and s["left"] > 0.72:
            x1 = min(x1, max(0.82, s["left"] - 0.012))
    if x1 - x0 < 0.55:
        x0, x1 = 0.055, 0.945
    if y1 - y0 < 0.50:
        y0, y1 = 0.12, 0.88
    return {"x": round(x0, 4), "y": round(y0, 4), "width": round(x1 - x0, 4), "height": round(y1 - y0, 4)}


def extract_images(prs: Presentation, assets_dir: Path):
    assets_dir.mkdir(parents=True, exist_ok=True)
    inventory = []
    seen_hashes = set()
    likely_logo = None
    for scope, shape in iter_shapes(prs):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        try:
            blob = shape.image.blob
            sha = hashlib.sha1(blob).hexdigest()
            if sha in seen_hashes:
                continue
            seen_hashes.add(sha)
            ext = (shape.image.ext or "png").lower()
            name = f"{scope}-{sha[:12]}.{ext}"
            path = assets_dir / name
            path.write_bytes(blob)
            area = max(1, shape.width * shape.height)
            inventory.append({"file": f"assets/{name}", "scope": scope, "sha1": sha, "area_emu2": area})
            if likely_logo is None and scope in {"master", "layout"}:
                likely_logo = f"assets/{name}"
        except Exception:
            continue
    return inventory, likely_logo


def source_fingerprint(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_html_layout(path: Path, kind: str, logo: str | None):
    logo_html = f'<img class="brand-logo" src="../{logo}" alt="logo">' if logo else ''
    title = {
        "cover": "Presentation title",
        "contents": "Contents",
        "section": "Section",
        "content": "Conclusion-led slide title",
        "ending": "Thank you",
    }[kind]
    body = {
        "cover": '<div class="cover-copy"><h1>Presentation title</h1><p>Subtitle / project / date</p></div>',
        "contents": '<main class="content-safe"><h1>Contents</h1><ol><li>Section one</li><li>Section two</li><li>Section three</li></ol></main>',
        "section": '<main class="content-safe section-center"><p class="eyebrow">01</p><h1>Section</h1><p>One-line section thesis</p></main>',
        "content": '<main class="content-safe"><h1>Conclusion-led slide title</h1><div class="content-canvas"><p>OpenDesign content canvas</p></div></main>',
        "ending": '<main class="content-safe section-center"><h1>Thank you</h1></main>',
    }[kind]
    html = f'''<!doctype html>
<html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><link rel="stylesheet" href="../brand.css"><title>{title}</title></head>
<body><section class="slide brand-shell {kind}">{logo_html}{body}<div class="brand-footer"></div></section></body></html>'''
    path.write_text(html, encoding="utf-8")


def render_references(source: Path, out_dir: Path, reps: dict[str, int | None], script_dir: Path):
    render_script = script_dir / "render_pptx.py"
    if not render_script.exists():
        return []
    tmp = out_dir / ".all-renders"
    try:
        proc = subprocess.run(["python", str(render_script), str(source), str(tmp)], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, check=False)
        if proc.returncode != 0:
            return []
        ref_dir = out_dir / "references"
        ref_dir.mkdir(parents=True, exist_ok=True)
        copied = []
        for role, idx in reps.items():
            if not idx:
                continue
            candidates = [tmp / f"slide-{idx:02d}.png", tmp / f"slide-{idx}.png"]
            src = next((p for p in candidates if p.exists()), None)
            if not src:
                continue
            dst = ref_dir / f"{role}.png"
            shutil.copy2(src, dst)
            copied.append(f"references/{dst.name}")
        return copied
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", type=Path)
    ap.add_argument("output_dir", type=Path)
    ap.add_argument("--name", help="Brand Pack display name")
    ap.add_argument("--no-render", action="store_true", help="Skip representative slide rendering")
    ns = ap.parse_args()

    source = ns.input.resolve()
    out = ns.output_dir.resolve()
    if not source.exists():
        raise SystemExit(f"Input not found: {source}")
    out.mkdir(parents=True, exist_ok=True)
    (out / "layouts").mkdir(exist_ok=True)
    (out / "assets").mkdir(exist_ok=True)
    (out / "references").mkdir(exist_ok=True)

    prs = Presentation(str(source))
    sw, sh = prs.slide_width, prs.slide_height
    reps = detect_representative_slides(prs)
    body_idxs = body_slide_indexes(len(prs.slides), reps)
    common = common_edge_shapes(prs, body_idxs)
    safe = estimate_safe_area(common)
    fonts, colors = collect_fonts_and_colors(prs)
    inventory, likely_logo = extract_images(prs, out / "assets")
    top_fonts = [name for name, _ in fonts.most_common(6)]
    top_colors = [color for color, _ in colors.most_common(8)]
    if not top_colors:
        top_colors = ["1F2937", "FFFFFF", "2563EB"]
    brand_name = ns.name or source.stem
    refs = [] if ns.no_render else render_references(source, out, reps, Path(__file__).resolve().parent)

    manifest = {
        "name": brand_name,
        "version": "1.0.0",
        "source": {"filename": source.name, "sha256": source_fingerprint(source)},
        "canvas": {
            "width_emu": sw,
            "height_emu": sh,
            "aspect_ratio": round(sw / sh, 6) if sh else None,
            "recommended_css_px": {"width": 1280, "height": round(1280 * sh / sw) if sw else 720},
        },
        "safe_area": safe,
        "colors": top_colors,
        "fonts": top_fonts,
        "locked_elements": common,
        "representative_slides": reps,
        "asset_inventory": inventory,
        "likely_logo": likely_logo,
        "reference_renders": refs,
        "notes": [
            "Automated extraction is a scaffold. Visually inspect representative renders and refine HTML/CSS shell before production use.",
            "This Brand Pack intentionally does not promise strict Slide Master/Layout preservation."
        ],
    }
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    primary = top_colors[0]
    secondary = top_colors[1] if len(top_colors) > 1 else "4B5563"
    bg = "FFFFFF" if "FFFFFF" in top_colors else (top_colors[2] if len(top_colors) > 2 else "FFFFFF")
    font = top_fonts[0] if top_fonts else "Microsoft YaHei"
    css = f'''/* Generated scaffold. Refine after visual inspection of references/*.png. */
:root {{
  --brand-primary: #{primary};
  --brand-secondary: #{secondary};
  --brand-bg: #{bg};
  --brand-text: #111827;
  --brand-font: "{font}", "Microsoft YaHei", "Noto Sans CJK SC", Arial, sans-serif;
  --safe-x: {safe['x'] * 100:.2f}%;
  --safe-y: {safe['y'] * 100:.2f}%;
  --safe-w: {safe['width'] * 100:.2f}%;
  --safe-h: {safe['height'] * 100:.2f}%;
}}
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; padding: 0; background: #ddd; }}
.slide {{ width: 1280px; height: {round(1280 * sh / sw) if sw else 720}px; position: relative; overflow: hidden; background: var(--brand-bg); color: var(--brand-text); font-family: var(--brand-font); }}
.content-safe {{ position: absolute; left: var(--safe-x); top: var(--safe-y); width: var(--safe-w); height: var(--safe-h); }}
.brand-logo {{ position: absolute; top: 3.2%; right: 4.2%; max-width: 12%; max-height: 8%; object-fit: contain; }}
.brand-footer {{ position: absolute; left: 5%; right: 5%; bottom: 2.5%; height: 1px; background: color-mix(in srgb, var(--brand-primary) 28%, transparent); }}
h1 {{ margin: 0 0 24px; font-size: 48px; line-height: 1.08; }}
p, li {{ font-size: 24px; line-height: 1.45; }}
.section-center {{ display: flex; flex-direction: column; justify-content: center; }}
.eyebrow {{ color: var(--brand-primary); font-weight: 700; letter-spacing: .08em; }}
.content-canvas {{ width: 100%; height: calc(100% - 90px); }}
'''
    (out / "brand.css").write_text(css, encoding="utf-8")

    md = f'''# {brand_name} Presentation Brand System

## Status

Generated from `{source.name}`. **Refine this file after visually inspecting `references/` renders.**

## Canvas

- Aspect ratio: {manifest['canvas']['aspect_ratio']}
- CSS authoring size: 1280×{manifest['canvas']['recommended_css_px']['height']}
- Safe content area: x={safe['x']:.3f}, y={safe['y']:.3f}, w={safe['width']:.3f}, h={safe['height']:.3f}

## Brand shell

Treat logo/background/header/footer/page-number chrome as locked when confirmed by the reference renders. Redesign the interior content canvas freely.

## Palette candidates

{chr(10).join(f'- `#{c}`' for c in top_colors)}

## Typography candidates

{chr(10).join(f'- {f}' for f in top_fonts) if top_fonts else '- Microsoft YaHei / system sans fallback'}

## Visual refinement checklist

- Confirm which extracted image is the real logo.
- Recreate any non-image master decoration in `brand.css`.
- Match header/footer line thickness and placement.
- Confirm safe content area does not collide with brand chrome.
- Preserve cover/contents/section/ending identity, while allowing body content to use richer layouts.
- Remove colors/assets that are accidental content rather than brand tokens.
'''
    (out / "BRAND_SYSTEM.md").write_text(md, encoding="utf-8")

    for kind in ["cover", "contents", "section", "content", "ending"]:
        write_html_layout(out / "layouts" / f"{kind}.html", kind, likely_logo)

    print(json.dumps({
        "brand_pack": str(out),
        "slides": len(prs.slides),
        "representative_slides": reps,
        "safe_area": safe,
        "fonts": top_fonts,
        "colors": top_colors,
        "assets": len(inventory),
        "reference_renders": refs,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

#!/usr/bin/env python3
"""Deterministic PPTX sanity checks.

Checks package readability, slide count, shape bounds, suspiciously tiny text, and empty slides.
Exits 2 when severe errors are found; warnings do not fail the command.
"""
from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path
from pptx import Presentation


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("input", type=Path)
    ap.add_argument("--json", dest="json_out", type=Path)
    ns = ap.parse_args()
    inp = ns.input.resolve()

    report = {"file": str(inp), "slides": 0, "errors": [], "warnings": [], "slide_reports": []}
    if not inp.exists():
        report["errors"].append("file_not_found")
    else:
        try:
            with zipfile.ZipFile(inp) as zf:
                bad = zf.testzip()
                if bad:
                    report["errors"].append(f"zip_crc_error:{bad}")
        except Exception as e:
            report["errors"].append(f"invalid_zip:{e}")

    if not report["errors"]:
        try:
            prs = Presentation(str(inp))
            sw, sh = prs.slide_width, prs.slide_height
            report["slides"] = len(prs.slides)
            if not prs.slides:
                report["errors"].append("no_slides")
            for si, slide in enumerate(prs.slides, 1):
                sr = {"slide": si, "shapes": len(slide.shapes), "errors": [], "warnings": []}
                if len(slide.shapes) == 0:
                    sr["warnings"].append("empty_slide")
                for shape in slide.shapes:
                    left, top, width, height = shape.left, shape.top, shape.width, shape.height
                    right, bottom = left + width, top + height
                    tolx = int(sw * 0.01)
                    toly = int(sh * 0.01)
                    if right < -tolx or bottom < -toly or left > sw + tolx or top > sh + toly:
                        sr["errors"].append(f"fully_outside:{shape.name}")
                    elif left < -tolx or top < -toly or right > sw + tolx or bottom > sh + toly:
                        sr["warnings"].append(f"partly_outside:{shape.name}")
                    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            for run in p.runs:
                                size = run.font.size
                                if size is not None and size.pt < 8:
                                    sr["warnings"].append(f"tiny_text:{shape.name}:{size.pt:.1f}pt")
                if sr["errors"]:
                    report["errors"].extend([f"slide_{si}:{e}" for e in sr["errors"]])
                if sr["warnings"]:
                    report["warnings"].extend([f"slide_{si}:{w}" for w in sr["warnings"]])
                report["slide_reports"].append(sr)
        except Exception as e:
            report["errors"].append(f"pptx_parse_error:{e}")

    text = json.dumps(report, ensure_ascii=False, indent=2)
    print(text)
    if ns.json_out:
        ns.json_out.parent.mkdir(parents=True, exist_ok=True)
        ns.json_out.write_text(text, encoding="utf-8")
    return 2 if report["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
